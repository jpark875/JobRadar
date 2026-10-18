import io
import re
from datetime import date
from pathlib import Path

from .skills import extract_skills

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1
)}
_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"


def _point(p: str) -> str:
    return rf"(?:(?P<{p}m>{_MONTH})\s+|(?P<{p}n>\d{{1,2}})/)?(?P<{p}y>(?:19|20)\d{{2}})"


_RANGE_RE = re.compile(
    _point("s") + r"\s*(?:-|–|—|to|until)\s*(?:" + _point("e") + r"|(?P<present>present|current|now|today))",
    re.IGNORECASE,
)
_EXPLICIT_RE = re.compile(r"(\d{1,2})\+?\s*(?:years|yrs)\s+(?:of\s+)?(?:professional\s+|industry\s+|work\s+)?experience", re.I)
_EDU_RE = re.compile(r"universit|college|school|bachelor|master|b\.s\.|b\.a\.|m\.s\.|ph\.?d|gpa|degree|coursework|academy", re.I)


class ResumeError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts).strip()
    if suffix in (".txt", ".md", ""):
        return data.decode("utf-8", errors="ignore").strip()
    raise ResumeError(f"Unsupported resume format: {suffix}. Use PDF, DOCX, or TXT.")


def _month_index(month: str | None, number: str | None, default: int) -> int:
    if month:
        return MONTHS.get(month[:3].lower(), default)
    if number and 1 <= int(number) <= 12:
        return int(number)
    return default


def estimate_years_experience(text: str, today: date | None = None) -> float:
    today = today or date.today()
    lines = text.splitlines()
    intervals = []
    for i, line in enumerate(lines):
        context = " ".join(lines[max(0, i - 1): i + 1])
        if _EDU_RE.search(context):
            continue
        for m in _RANGE_RE.finditer(line):
            start = int(m.group("sy")) * 12 + _month_index(m.group("sm"), m.group("sn"), 1) - 1
            if m.group("present"):
                end = today.year * 12 + today.month - 1
            else:
                end = int(m.group("ey")) * 12 + _month_index(m.group("em"), m.group("en"), 12) - 1
            if start <= end <= today.year * 12 + 12:
                intervals.append((start, end))
    months = 0
    current = None
    for start, end in sorted(intervals):
        if current and start <= current[1]:
            current = (current[0], max(current[1], end))
        else:
            if current:
                months += current[1] - current[0] + 1
            current = (start, end)
    if current:
        months += current[1] - current[0] + 1
    from_ranges = round(months / 12, 1)
    explicit = [int(x) for x in _EXPLICIT_RE.findall(text)]
    return max([from_ranges, *explicit]) if explicit else from_ranges


def analyze(text: str) -> dict:
    return {"skills": extract_skills(text), "years_experience": estimate_years_experience(text)}
